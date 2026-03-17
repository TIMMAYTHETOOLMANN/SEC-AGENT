"""
JLAW Presentation Deck Builder — PPTX with Embedded Charts
Generates a professional forensic intelligence presentation deck using python-pptx.

Slide structure:
  1. Title slide
  2. Executive summary with key metrics
  3. Insider Trading Timeline chart
  4. Anomaly Severity Heatmap chart
  5. 13D Staleness Gauge chart
  6. Filing Timeliness Box Plot chart
  7. Compounding Pattern Network chart
  8. Regulatory Scorecard Radar chart
  9. Key findings summary table
  10. Recommended actions

Usage:
    from src.visualization.deck_builder import DeckBuilder
    builder = DeckBuilder(anomaly_db, chart_images)
    builder.build("data/visualizations/JLAW_Forensic_Deck.pptx")
"""

import json
from pathlib import Path
from datetime import datetime
from typing import Dict, List, Optional

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE


# ═══════════════════════════════════════════════
# THEME COLORS (matching JLAW palette)
# ═══════════════════════════════════════════════

BG_COLOR      = RGBColor(0x0d, 0x11, 0x17)
PANEL_COLOR   = RGBColor(0x16, 0x1b, 0x22)
TEXT_COLOR     = RGBColor(0xc9, 0xd1, 0xd9)
TEXT_DIM       = RGBColor(0x8b, 0x94, 0x9e)
ACCENT_COLOR   = RGBColor(0x58, 0xa6, 0xff)
DANGER_COLOR   = RGBColor(0xf8, 0x51, 0x49)
WARN_COLOR     = RGBColor(0xd2, 0x99, 0x22)
GREEN_COLOR    = RGBColor(0x3f, 0xb9, 0x50)
WHITE_COLOR    = RGBColor(0xff, 0xff, 0xff)

SEVERITY_PPT_COLORS = {
    "STRUCTURAL": DANGER_COLOR,
    "HIGH":       RGBColor(0xf0, 0x88, 0x3e),
    "MEDIUM":     WARN_COLOR,
    "LOW":        ACCENT_COLOR,
}


def _set_slide_bg(slide, color: RGBColor = BG_COLOR):
    """Set slide background to dark theme."""
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color


def _add_text_box(slide, left, top, width, height, text, font_size=12,
                  color=TEXT_COLOR, bold=False, alignment=PP_ALIGN.LEFT,
                  font_name="Segoe UI"):
    """Add a formatted text box to a slide."""
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = font_name
    p.alignment = alignment

    return txBox


class DeckBuilder:
    """Builds a PPTX presentation deck with embedded forensic charts."""

    def __init__(self, anomaly_db: Dict, chart_images: Dict[str, Path] = None):
        self.db = anomaly_db
        self.anomalies = anomaly_db.get("anomalies", [])
        self.patterns = anomaly_db.get("patterns", [])
        self.chart_images = chart_images or {}
        self.prs = Presentation()
        self.prs.slide_width = Inches(13.333)  # Widescreen 16:9
        self.prs.slide_height = Inches(7.5)

    def build(self, output_path: Path) -> Path:
        """Build the complete presentation deck."""
        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)

        self._add_title_slide()
        self._add_executive_summary()
        self._add_chart_slides()
        self._add_findings_table()
        self._add_recommendations()

        self.prs.save(output_path)
        return output_path

    def _add_title_slide(self):
        """Slide 1: Title."""
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])  # Blank layout
        _set_slide_bg(slide)

        # Classification header
        _add_text_box(slide, Inches(0.5), Inches(0.3), Inches(12), Inches(0.4),
                      "PRIVILEGED & CONFIDENTIAL — ATTORNEY WORK PRODUCT",
                      font_size=10, color=DANGER_COLOR, bold=True,
                      alignment=PP_ALIGN.CENTER)

        # Main title
        _add_text_box(slide, Inches(1), Inches(2.2), Inches(11), Inches(1.2),
                      "JLAW FORENSIC INTELLIGENCE DASHBOARD",
                      font_size=36, color=ACCENT_COLOR, bold=True,
                      alignment=PP_ALIGN.CENTER)

        # Subtitle
        _add_text_box(slide, Inches(1), Inches(3.5), Inches(11), Inches(0.6),
                      "Nike, Inc. (CIK 0000320187) — Seven-Year Investigation",
                      font_size=18, color=TEXT_COLOR, alignment=PP_ALIGN.CENTER)

        _add_text_box(slide, Inches(1), Inches(4.2), Inches(11), Inches(0.5),
                      "FY2019 through Q1 CY2026",
                      font_size=14, color=TEXT_DIM, alignment=PP_ALIGN.CENTER)

        # Date
        now = datetime.now().strftime("%B %d, %Y")
        _add_text_box(slide, Inches(1), Inches(5.5), Inches(11), Inches(0.4),
                      f"Generated: {now} | JLAW Agent Platform v5.0",
                      font_size=11, color=TEXT_DIM, alignment=PP_ALIGN.CENTER)

    def _add_executive_summary(self):
        """Slide 2: Executive summary with key metrics."""
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        _set_slide_bg(slide)

        _add_text_box(slide, Inches(0.5), Inches(0.3), Inches(12), Inches(0.5),
                      "EXECUTIVE SUMMARY", font_size=24, color=ACCENT_COLOR, bold=True)

        # Metric cards
        structural = sum(1 for a in self.anomalies if a.get("severity") == "STRUCTURAL")
        high = sum(1 for a in self.anomalies if a.get("severity") == "HIGH")
        medium = sum(1 for a in self.anomalies if a.get("severity") == "MEDIUM")

        metrics = [
            (str(len(self.anomalies)), "Total\nAnomalies", ACCENT_COLOR),
            (str(structural), "Structural\nFailures", DANGER_COLOR),
            (str(high), "High\nSeverity", RGBColor(0xf0, 0x88, 0x3e)),
            (str(len(self.patterns)), "Compounding\nPatterns", WARN_COLOR),
            ("7", "Fiscal\nYears", GREEN_COLOR),
            ("$565M+", "Cumulative\nInsider Sales", DANGER_COLOR),
        ]

        card_width = Inches(1.8)
        card_height = Inches(1.5)
        start_x = Inches(0.5)
        gap = Inches(0.25)

        for i, (value, label, color) in enumerate(metrics):
            x = start_x + i * (card_width + gap)
            y = Inches(1.2)

            # Card background
            shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                           x, y, card_width, card_height)
            shape.fill.solid()
            shape.fill.fore_color.rgb = PANEL_COLOR
            shape.line.color.rgb = RGBColor(0x21, 0x26, 0x2d)
            shape.line.width = Pt(1)

            # Value
            _add_text_box(slide, x + Inches(0.1), y + Inches(0.15),
                         card_width - Inches(0.2), Inches(0.7),
                         value, font_size=28, color=color, bold=True,
                         alignment=PP_ALIGN.CENTER)

            # Label
            _add_text_box(slide, x + Inches(0.1), y + Inches(0.85),
                         card_width - Inches(0.2), Inches(0.5),
                         label, font_size=10, color=TEXT_DIM,
                         alignment=PP_ALIGN.CENTER)

        # Key findings text
        findings_text = (
            "This dashboard synthesizes findings from a seven-year forensic investigation "
            "of Nike, Inc. spanning FY2019 through Q1 CY2026.\n\n"
            "Critical findings include:\n"
            "• Swoosh LLC 13D unamended since June 2016 — active Rule 13d-2 violation\n"
            "• Parker cumulative insider sales exceeding $565M under self-approval pathway\n"
            "• Travis Knight late Form 4 (Jan 2026) — first Knight family §16 failure\n"
            "• 14-year SEC correspondence gap (no CORRESP filings since ~2011)\n"
            "• Buy/sell divergence: $6.2M director purchases vs $57.5M Parker sales (FY2025)"
        )
        _add_text_box(slide, Inches(0.5), Inches(3.2), Inches(12), Inches(3.5),
                      findings_text, font_size=13, color=TEXT_COLOR)

    def _add_chart_slides(self):
        """Slides 3-8: One chart per slide."""
        chart_config = [
            ("insider_timeline",    "INSIDER TRADING TIMELINE",
             "Seven-year insider selling cadence with cumulative overlay. "
             "Parker quarterly sales pattern + FY2025 buy/sell divergence."),
            ("severity_heatmap",    "ANOMALY SEVERITY HEATMAP",
             "Year × Category matrix showing weighted severity scores. "
             "Darker cells indicate higher concentration of severe anomalies."),
            ("staleness_gauge",     "SWOOSH 13D STALENESS GAUGE",
             "Swoosh LLC Schedule 13D has not been amended since June 2016 — "
             "an active Rule 13d-2 violation spanning 9.7+ years."),
            ("timeliness_boxplot",  "FORM 4 FILING TIMELINESS",
             "Distribution of business days to file Form 4 per insider. "
             "SEC Rule 16a-3(g) requires filing within 2 business days."),
            ("pattern_network",     "COMPOUNDING PATTERN NETWORK",
             "Force-directed graph linking anomalies to their compounding patterns. "
             "Node color indicates severity; size indicates linkage density."),
            ("regulatory_radar",    "REGULATORY ENFORCEMENT SCORECARD",
             "Multi-axis readiness assessment across SEC, DOJ, ISS, Congressional, "
             "PCAOB, and whistleblower enforcement dimensions."),
        ]

        for chart_key, title, description in chart_config:
            slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
            _set_slide_bg(slide)

            # Title
            _add_text_box(slide, Inches(0.5), Inches(0.3), Inches(12), Inches(0.5),
                          title, font_size=22, color=ACCENT_COLOR, bold=True)

            # Description
            _add_text_box(slide, Inches(0.5), Inches(0.85), Inches(12), Inches(0.5),
                          description, font_size=11, color=TEXT_DIM)

            # Chart image
            img_path = self.chart_images.get(chart_key)
            if img_path and Path(img_path).exists():
                # Center the image
                img_width = Inches(11)
                img_height = Inches(5.5)
                img_left = Inches(1.15)
                img_top = Inches(1.5)
                slide.shapes.add_picture(str(img_path), img_left, img_top,
                                        img_width, img_height)
            else:
                # Placeholder
                _add_text_box(slide, Inches(3), Inches(3.5), Inches(7), Inches(1),
                              f"[Chart image not available: {chart_key}]",
                              font_size=14, color=TEXT_DIM, alignment=PP_ALIGN.CENTER)

    def _add_findings_table(self):
        """Slide 9: Key findings summary table."""
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        _set_slide_bg(slide)

        _add_text_box(slide, Inches(0.5), Inches(0.3), Inches(12), Inches(0.5),
                      "KEY FINDINGS CATALOG", font_size=22, color=ACCENT_COLOR, bold=True)

        # Build table from top anomalies
        top_anomalies = sorted(
            self.anomalies,
            key=lambda a: {"STRUCTURAL": 0, "HIGH": 1, "MEDIUM": 2, "LOW": 3}.get(
                a.get("severity", "LOW"), 4)
        )[:12]  # Top 12

        if not top_anomalies:
            _add_text_box(slide, Inches(0.5), Inches(1.5), Inches(12), Inches(1),
                          "No anomalies available. Run Phase 2 (cross-reference) first.",
                          font_size=14, color=TEXT_DIM)
            return

        rows = len(top_anomalies) + 1
        cols = 4
        table = slide.shapes.add_table(rows, cols,
                                        Inches(0.5), Inches(1.2),
                                        Inches(12), Inches(5.5)).table

        # Column widths
        table.columns[0].width = Inches(1.5)
        table.columns[1].width = Inches(1.2)
        table.columns[2].width = Inches(1.3)
        table.columns[3].width = Inches(8.0)

        # Header row
        headers = ["ID", "Severity", "Fiscal Year", "Finding"]
        for j, header in enumerate(headers):
            cell = table.cell(0, j)
            cell.text = header
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor(0x21, 0x26, 0x2d)
            for p in cell.text_frame.paragraphs:
                p.font.size = Pt(10)
                p.font.bold = True
                p.font.color.rgb = ACCENT_COLOR

        # Data rows
        for i, a in enumerate(top_anomalies):
            row_idx = i + 1
            values = [
                a.get("anomaly_id", ""),
                a.get("severity", ""),
                a.get("fiscal_year", ""),
                (a.get("title", "") or "")[:90],
            ]
            for j, val in enumerate(values):
                cell = table.cell(row_idx, j)
                cell.text = str(val)
                cell.fill.solid()
                cell.fill.fore_color.rgb = PANEL_COLOR if i % 2 == 0 else BG_COLOR
                for p in cell.text_frame.paragraphs:
                    p.font.size = Pt(9)
                    p.font.color.rgb = TEXT_COLOR
                    if j == 1:  # Severity column
                        p.font.color.rgb = SEVERITY_PPT_COLORS.get(val, TEXT_COLOR)
                        p.font.bold = True

    def _add_recommendations(self):
        """Slide 10: Recommended regulatory actions."""
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        _set_slide_bg(slide)

        _add_text_box(slide, Inches(0.5), Inches(0.3), Inches(12), Inches(0.5),
                      "RECOMMENDED REGULATORY ACTIONS", font_size=22,
                      color=ACCENT_COLOR, bold=True)

        recommendations = [
            ("SEC Division of Enforcement", [
                "Formal investigation of insider trading patterns under Section 10(b)",
                "Compelled amendment of Swoosh LLC stale Schedule 13D",
                "Enforcement action for Section 16(a) filing delinquencies",
                "Review of Exhibit 19 pre-clearance self-approval mechanism",
            ]),
            ("DOJ Fraud Section", [
                "Criminal referral for securities fraud pattern ($565M+ cumulative)",
                "Review of MNPI timing around CEO termination sequence",
            ]),
            ("ISS / Congressional", [
                "Governance engagement on dual-class voting concentration",
                "Legislative reform for mandatory annual 13D updates",
                "Enhanced pre-clearance audit requirements",
            ]),
        ]

        y = Inches(1.2)
        for agency, actions in recommendations:
            _add_text_box(slide, Inches(0.5), y, Inches(12), Inches(0.35),
                          agency, font_size=14, color=WARN_COLOR, bold=True)
            y += Inches(0.4)

            for action in actions:
                _add_text_box(slide, Inches(0.8), y, Inches(11.5), Inches(0.3),
                              f"→  {action}", font_size=11, color=TEXT_COLOR)
                y += Inches(0.3)

            y += Inches(0.2)

        # Footer
        _add_text_box(slide, Inches(0.5), Inches(6.8), Inches(12), Inches(0.4),
                      "JLAW Agent Platform v5.0 | All findings cite exact EDGAR accession numbers",
                      font_size=10, color=TEXT_DIM, alignment=PP_ALIGN.CENTER)
