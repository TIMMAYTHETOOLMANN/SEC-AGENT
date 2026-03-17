"""
JLAW XLS/XLSX Filing Parser
Parses Excel spreadsheet downloads from SEC EDGAR into structured JSON.
Classifies sheets as Form 4 transactions, beneficial ownership tables,
vote tallies, or compensation summaries.
"""

import json
import re
from pathlib import Path
from datetime import datetime
from typing import Dict, List, Optional

try:
    import openpyxl
except ImportError:
    openpyxl = None

try:
    import xlrd
except ImportError:
    xlrd = None


# Sheet classification keywords
SHEET_CLASSIFIERS = {
    "form4_transactions": [
        "transaction", "form 4", "insider", "shares acquired",
        "shares disposed", "exercise price", "10b5-1",
    ],
    "beneficial_ownership": [
        "beneficial", "ownership", "class a", "class b",
        "percent of class", "sole voting", "shared voting",
    ],
    "vote_tally": [
        "votes for", "votes against", "abstain", "broker non-vote",
        "withheld", "proposal", "say-on-pay", "director election",
    ],
    "compensation": [
        "salary", "bonus", "stock awards", "option awards",
        "non-equity incentive", "all other compensation", "total",
        "summary compensation", "NEO",
    ],
    "exhibit_index": [
        "exhibit", "description", "filed herewith",
        "incorporated by reference", "accession",
    ],
}


def classify_sheet(headers: List[str], sample_data: List[List]) -> str:
    """Classify a sheet based on its headers and sample data."""
    all_text = " ".join(str(h).lower() for h in headers if h)
    # Also check first few rows
    for row in sample_data[:5]:
        all_text += " " + " ".join(str(c).lower() for c in row if c)

    scores = {}
    for category, keywords in SHEET_CLASSIFIERS.items():
        score = sum(1 for kw in keywords if kw in all_text)
        if score > 0:
            scores[category] = score

    if scores:
        return max(scores, key=scores.get)
    return "unknown"


def parse_xlsx(file_path: Path) -> Dict:
    """Parse an XLSX file into structured JSON."""
    if openpyxl is None:
        return {"error": "openpyxl not installed", "source": str(file_path)}

    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    sheets = {}

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = []
        headers = []

        for i, row in enumerate(ws.iter_rows(values_only=True)):
            clean_row = [str(c).strip() if c is not None else "" for c in row]
            if i == 0:
                headers = clean_row
            rows.append(clean_row)

        if not rows:
            continue

        # Classify the sheet
        classification = classify_sheet(headers, rows[1:6])

        # Extract structured data based on classification
        sheet_data = {
            "sheet_name": sheet_name,
            "classification": classification,
            "headers": headers,
            "row_count": len(rows) - 1,
            "data": [],
        }

        if classification == "form4_transactions":
            sheet_data["data"] = _extract_form4_data(headers, rows[1:])
        elif classification == "beneficial_ownership":
            sheet_data["data"] = _extract_ownership_data(headers, rows[1:])
        elif classification == "vote_tally":
            sheet_data["data"] = _extract_vote_data(headers, rows[1:])
        elif classification == "compensation":
            sheet_data["data"] = _extract_comp_data(headers, rows[1:])
        else:
            # Raw extraction — first 100 rows
            sheet_data["data"] = [
                dict(zip(headers, row)) for row in rows[1:101]
                if any(c.strip() for c in row)
            ]

        sheets[sheet_name] = sheet_data

    wb.close()

    return {
        "source_file": str(file_path),
        "file_type": "xlsx",
        "sheet_count": len(sheets),
        "sheets": sheets,
        "parsed_at": datetime.now().isoformat(),
    }


def parse_xls(file_path: Path) -> Dict:
    """Parse a legacy XLS file."""
    if xlrd is None:
        return {"error": "xlrd not installed. Run: pip install xlrd>=2.0.0", "source": str(file_path)}

    wb = xlrd.open_workbook(str(file_path))
    sheets = {}

    for sheet_name in wb.sheet_names():
        ws = wb.sheet_by_name(sheet_name)
        rows = []
        headers = []

        for i in range(ws.nrows):
            row = [str(ws.cell_value(i, j)).strip() for j in range(ws.ncols)]
            if i == 0:
                headers = row
            rows.append(row)

        if not rows:
            continue

        classification = classify_sheet(headers, rows[1:6])

        sheet_data = {
            "sheet_name": sheet_name,
            "classification": classification,
            "headers": headers,
            "row_count": len(rows) - 1,
            "data": [
                dict(zip(headers, row)) for row in rows[1:101]
                if any(c.strip() for c in row)
            ],
        }

        sheets[sheet_name] = sheet_data

    return {
        "source_file": str(file_path),
        "file_type": "xls",
        "sheet_count": len(sheets),
        "sheets": sheets,
        "parsed_at": datetime.now().isoformat(),
    }


def _extract_form4_data(headers: List[str], rows: List[List]) -> List[Dict]:
    """Extract Form 4 transaction data."""
    transactions = []
    h_lower = [h.lower() for h in headers]

    def _find_col(keywords):
        for kw in keywords:
            for i, h in enumerate(h_lower):
                if kw in h:
                    return i
        return None

    date_col = _find_col(["date", "transaction date"])
    name_col = _find_col(["name", "insider", "filer", "reporting person"])
    code_col = _find_col(["code", "transaction code", "type"])
    shares_col = _find_col(["shares", "amount", "quantity"])
    price_col = _find_col(["price", "exercise"])

    for row in rows:
        if not any(c.strip() for c in row):
            continue
        txn = {}
        if date_col is not None and date_col < len(row):
            txn["date"] = row[date_col]
        if name_col is not None and name_col < len(row):
            txn["name"] = row[name_col]
        if code_col is not None and code_col < len(row):
            txn["code"] = row[code_col]
        if shares_col is not None and shares_col < len(row):
            txn["shares"] = row[shares_col]
        if price_col is not None and price_col < len(row):
            txn["price"] = row[price_col]
        if txn:
            transactions.append(txn)

    return transactions


def _extract_ownership_data(headers: List[str], rows: List[List]) -> List[Dict]:
    """Extract beneficial ownership table data."""
    owners = []
    for row in rows:
        if not any(c.strip() for c in row):
            continue
        entry = dict(zip(headers, row))
        # Clean up numeric fields
        for key in entry:
            val = entry[key]
            if isinstance(val, str):
                cleaned = re.sub(r'[,$%]', '', val).strip()
                try:
                    entry[key] = float(cleaned) if '.' in cleaned else int(cleaned)
                except (ValueError, TypeError):
                    pass
        owners.append(entry)
    return owners


def _extract_vote_data(headers: List[str], rows: List[List]) -> List[Dict]:
    """Extract annual meeting vote tally data."""
    votes = []
    for row in rows:
        if not any(c.strip() for c in row):
            continue
        entry = dict(zip(headers, row))
        votes.append(entry)
    return votes


def _extract_comp_data(headers: List[str], rows: List[List]) -> List[Dict]:
    """Extract compensation summary data."""
    comp = []
    for row in rows:
        if not any(c.strip() for c in row):
            continue
        entry = dict(zip(headers, row))
        # Try to parse monetary values
        for key in entry:
            val = entry[key]
            if isinstance(val, str):
                cleaned = re.sub(r'[$,]', '', val).strip()
                try:
                    entry[key] = float(cleaned)
                except (ValueError, TypeError):
                    pass
        comp.append(entry)
    return comp


def parse_spreadsheet(file_path: Path) -> Dict:
    """Parse any spreadsheet file (XLS or XLSX)."""
    suffix = file_path.suffix.lower()
    if suffix == ".xlsx" or suffix == ".xlsm":
        return parse_xlsx(file_path)
    elif suffix == ".xls":
        return parse_xls(file_path)
    else:
        return {"error": f"Unsupported format: {suffix}", "source": str(file_path)}


def parse_spreadsheet_directory(directory: Path, output_dir: Path) -> List[Dict]:
    """Parse all spreadsheets in a directory tree."""
    output_dir.mkdir(parents=True, exist_ok=True)
    results = []

    for xls_file in sorted(directory.rglob("*.xls*")):
        if xls_file.name.startswith("~"):
            continue  # Skip temp files
        try:
            parsed = parse_spreadsheet(xls_file)
            out_name = xls_file.stem + "_parsed.json"
            out_path = output_dir / out_name
            with open(out_path, "w") as f:
                json.dump(parsed, f, indent=2, default=str)
            results.append({
                "file": str(xls_file),
                "sheets": parsed.get("sheet_count", 0),
                "classifications": {
                    s["sheet_name"]: s["classification"]
                    for s in parsed.get("sheets", {}).values()
                } if isinstance(parsed.get("sheets"), dict) else {},
                "status": "OK",
            })
        except Exception as e:
            results.append({"file": str(xls_file), "status": "ERROR", "error": str(e)})

    return results


if __name__ == "__main__":
    import sys
    if len(sys.argv) < 2:
        print("Usage: python -m src.ingestion.xls_parser <file_or_directory> [output_dir]")
        sys.exit(1)

    path = Path(sys.argv[1])
    if path.is_file():
        result = parse_spreadsheet(path)
        print(json.dumps(result, indent=2, default=str))
    elif path.is_dir():
        out = Path(sys.argv[2]) if len(sys.argv) > 2 else path / "parsed"
        results = parse_spreadsheet_directory(path, out)
        for r in results:
            print(f"  [{r['status']}] {r['file']} — {r.get('sheets', '?')} sheets")
